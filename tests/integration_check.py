"""End-to-end integration check for Tasks & Meetings + cross-system consistency.

Runs against a THROWAWAY temp database (config.DATABASE_PATH is repointed before
any DB call), so the real ./data/yordamchi.db is never touched.

Run:  venv/bin/python tests/integration_check.py
"""
import asyncio
import os
import sys
from datetime import datetime, timedelta, time as dtime

# Allow running from anywhere — project root holds config.py / database.py.
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Repoint DB to a temp file BEFORE importing database/handlers.
import config
_TMP = "/tmp/yordamchi_integration_test.db"
if os.path.exists(_TMP):
    os.remove(_TMP)
config.DATABASE_PATH = _TMP

import database  # noqa: E402
import handlers  # noqa: E402

TZ = database.TZ
PASS = 0
FAIL = 0
FAILED = []


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        FAILED.append(name)
        print(f"  ❌ {name}   {detail}")


def has(lst, _id):
    return any((x.get("id") == _id) for x in lst)


async def main():
    await database.init()
    now = datetime.now(TZ)
    today_start = TZ.localize(datetime.combine(now.date(), dtime.min))
    wk_start = today_start.isoformat()
    wk_end = (today_start + timedelta(days=7)).isoformat()

    print("\n── A. VAZIFALAR ──")
    tid = await database.create_task({"title": "TEST vazifa A", "priority": "P2",
                                      "status": "todo", "assignee": "Aziz",
                                      "deadline": (now + timedelta(days=1)).isoformat(),
                                      "tags": ["test", "muhim"]})
    t = await database.get_task(tid)
    check("A1 create_task saqlanadi", t is not None and t["title"] == "TEST vazifa A")
    check("A1 tags ro'yxat sifatida o'qiladi", isinstance(t.get("tags"), list) and "test" in t["tags"],
          f"tags={t.get('tags')!r}")
    active = await database.list_tasks(status_in=["todo", "in_progress"], limit=200)
    done = await database.list_tasks(status_in=["done"], limit=200)
    check("A2 yangi vazifa AKTIV filtrda", has(active, tid))
    check("A2 yangi vazifa BAJARILGAN filtrda EMAS", not has(done, tid))

    # cross-system snapshot before complete
    active_before = len(active)
    done_before = len(done)

    print("\n── C1. todo → done (cross-system) ──")
    ok = await database.complete_task(tid)
    t = await database.get_task(tid)
    active2 = await database.list_tasks(status_in=["todo", "in_progress"], limit=200)
    done2 = await database.list_tasks(status_in=["done"], limit=200)
    done_today = await database.list_tasks_done_today()
    check("C1 complete_task ok", ok and t["status"] == "done")
    check("C1 AKTIV'dan yo'qoldi", not has(active2, tid))
    check("C1 BAJARILGAN'da paydo bo'ldi", has(done2, tid))
    check("C1 'bugun bajarilgan'da bor", has(done_today, tid))
    check("C1 sanoq izchil (aktiv -1, done +1)",
          len(active2) == active_before - 1 and len(done2) == done_before + 1,
          f"aktiv {active_before}->{len(active2)}, done {done_before}->{len(done2)}")

    print("\n── A4. Tahrir (prioritet) ──")
    await database.update_task(tid, {"priority": "P0"}, source="edit")
    t = await database.get_task(tid)
    check("A4 prioritet P0 ga o'zgardi", t["priority"] == "P0")

    print("\n── C2. O'chirish (cross-system) ──")
    await database.delete_task(tid)
    gone = await database.get_task(tid)
    a3 = await database.list_tasks(status_in=["todo", "in_progress"], limit=200)
    d3 = await database.list_tasks(status_in=["done"], limit=200)
    allt = await database.list_tasks(limit=500)
    check("C2 get_task None", gone is None)
    check("C2 hech bir filtrda yo'q (aktiv/done/barchasi)",
          not has(a3, tid) and not has(d3, tid) and not has(allt, tid))

    print("\n── A6. Bugungi muddatli vazifa ──")
    tid2 = await database.create_task({"title": "TEST bugun", "priority": "P1", "status": "todo",
                                       "deadline": (today_start + timedelta(hours=15)).isoformat()})
    today_tasks = await database.list_today_tasks()
    check("A6 bugungi muddat → list_today_tasks", has(today_tasks, tid2))
    await database.delete_task(tid2)

    print("\n── B/D1. Uchrashuv + AGENDA RO'YXATI (regressiya) ──")
    mdata = {"title": "TEST uchrashuv", "datetime_start": (now + timedelta(days=1)).isoformat(),
             "datetime_end": (now + timedelta(days=1, hours=1)).isoformat(),
             "participants": ["Dinislam", "Aziz"],
             "agenda": ["Byudjet", "Marketing", "Keyingi qadamlar"]}
    mid = await database.create_meeting(mdata)
    m = await database.get_meeting(mid)
    check("D1 agenda=ro'yxat bo'lsa ham SAQLANADI", m is not None, "create_meeting yiqildi")
    check("D1 agenda matn sifatida o'qiladi", isinstance(m.get("agenda"), str) and "Byudjet" in m["agenda"],
          f"agenda={m.get('agenda')!r}")
    check("D1 participants ro'yxat sifatida o'qiladi", isinstance(m.get("participants"), list) and "Aziz" in m["participants"])

    print("\n── B2/D2. Bugun ERTAROQ uchrashuv → 'Haftalik'da (regressiya) ──")
    earlier = today_start + timedelta(minutes=1)  # bugun, lekin hozirdan oldin
    mid2 = await database.create_meeting({"title": "TEST ertaroq", "datetime_start": earlier.isoformat(),
                                          "datetime_end": (earlier + timedelta(hours=1)).isoformat(),
                                          "participants": ["X"], "agenda": ["a"]})
    week = await database.list_meetings_in_window(wk_start, wk_end)
    week_now = await database.list_meetings_in_window(now.isoformat(), wk_end)  # eski (buggy) oyna
    today_m = await database.list_today_meetings()
    check("D2 'Haftalik' (bugun boshidan) ko'rsatadi", has(week, mid2))
    check("D2 'Bugun' filtrida bor", has(today_m, mid2))
    check("D2 (eslatma) eski now-oyna ko'rsatMASdi", not has(week_now, mid2),
          "eski oyna ham ko'rsatsa, regressiya emas edi")

    print("\n── C4. Reschedule (cross-system) ──")
    new_start = (now + timedelta(days=2, hours=3)).isoformat()
    await database.update_meeting(mid, {"datetime_start": new_start})
    m = await database.get_meeting(mid)
    check("C4 datetime_start yangilandi", m["datetime_start"] == new_start)
    w2 = await database.list_meetings_in_window(wk_start, (today_start + timedelta(days=10)).isoformat())
    found = [x for x in w2 if x["id"] == mid]
    check("C4 oyna so'rovida yangi vaqt bilan chiqadi", bool(found) and found[0]["datetime_start"] == new_start)

    print("\n── C5. Bekor (cross-system) ──")
    await database.cancel_meeting(mid)
    await database.cancel_meeting(mid2)
    check("C5 get_meeting None", (await database.get_meeting(mid)) is None)
    w3 = await database.list_meetings_in_window(wk_start, (today_start + timedelta(days=10)).isoformat())
    check("C5 oyna so'rovidan yo'qoldi", not has(w3, mid) and not has(w3, mid2))

    print("\n── D5. Sana parser — sabab ajratiladi ──")
    cases = {"2 soat": None, "ertaga 09:00": None, "999 soat": "too_far",
             "2026-13-45 10:00": "invalid", "falongdek": "unparsable"}
    for txt, exp in cases.items():
        iso, reason = await handlers._parse_deadline_natural(txt)
        if exp is None:
            check(f"D5 '{txt}' → parse OK", iso is not None and reason is None, f"reason={reason}")
        else:
            check(f"D5 '{txt}' → '{exp}'", iso is None and reason == exp, f"reason={reason}")

    print("\n── Stats ishlaydi (cross-system yuzasi) ──")
    try:
        st = await database.executive_stats(days=7)
        check("executive_stats xatosiz ishlaydi", isinstance(st, dict) and "tasks" in st)
    except Exception as e:
        check("executive_stats xatosiz ishlaydi", False, f"{type(e).__name__}: {e}")

    print("\n── Export / Import (Excel) ──")
    import io as _io
    from openpyxl import load_workbook as _lwb

    class _ExpMsg:
        chat = type("C", (), {"id": 1})()
        text = "/export"
        captured = {}
        async def answer_document(self, file, caption=None, parse_mode=None):
            _ExpMsg.captured["b"] = file.data
        async def answer(self, *a, **k):
            _ExpMsg.captured["t"] = a

    await database.create_task({"title": "Export sinov", "assignee": "J.K", "priority": "P0",
                                "status": "todo", "deadline": (datetime.now(TZ) + timedelta(days=2)).isoformat(),
                                "description": "izoh"})
    await handlers.cmd_export(_ExpMsg())
    try:
        _ws = _lwb(_io.BytesIO(_ExpMsg.captured["b"])).active
        check("Export: xlsx yaratildi", _ExpMsg.captured.get("b") is not None)
        check("Export: sarlavha B1", _ws["B1"].value == "VAZIFALAR RO'YXATI")
        check("Export: header A3=№ B3=Vazifa", _ws["A3"].value == "№" and _ws["B3"].value == "Vazifa")
        check("Export: header Arial 14", _ws["B3"].font.name == "Arial" and int(_ws["B3"].font.sz) == 14)
        check("Export: Kategoriya ustuni (H)", _ws["H3"].value == "Kategoriya")
        check("Export: yashirin ID ustuni (I)", _ws["I3"].value == "ID" and bool(_ws.column_dimensions["I"].hidden))
    except Exception as e:
        check("Export: xlsx tahlili", False, f"{type(e).__name__}: {e}")

    _acts = handlers._structured_tasks_from_table(
        [("№", "Vazifa", "Ijrochi", "Muddat", "Ustuvorlik", "Holat", "Izoh"),
         (1, "Import A", "Aziz", "08-06-2026", "Shoshilinch", "Aktiv", "x")])
    check("Import (struktura): 1 vazifa", len(_acts) == 1 and _acts[0]["data"]["title"] == "Import A")
    check("Import: Shoshilinch→P0", bool(_acts) and _acts[0]["data"]["priority"] == "P0")
    check("Import: noma'lum ustun → [] (aqlli yo'lga)", handlers._structured_tasks_from_table([("Mahsulot", "Narx", "Soni"), ("Olma", "5000", "10")]) == [])
    check("Import: Mas'ul ustuni → ijrochi", (lambda a: bool(a) and a[0]["data"].get("assignee") == "O.X")(handlers._structured_tasks_from_table([("Vazifa", "Mas'ul"), ("Ish", "O.X")])))
    check("Import: izoh==ijrochi → izoh tushiriladi", (lambda a: bool(a) and not a[0]["data"].get("description"))(handlers._structured_tasks_from_table([("Vazifa", "Ijrochi", "Izoh"), ("Ish", "Aziz", "Aziz")])))

    print("\n── Kategoriyalar ──")
    _cid = await database.create_task({"title": "Kat sinov", "priority": "P1", "status": "todo", "category": "Shartnomalar"})
    _ct = await database.get_task(_cid)
    check("Kategoriya saqlanadi", _ct.get("category") == "Shartnomalar")
    _cats = await database.list_task_categories()
    check("list_task_categories ishlaydi", any(c["category"] == "Shartnomalar" and c["count"] >= 1 for c in _cats))
    check("list_tasks_by_category", any(t["id"] == _cid for t in await database.list_tasks_by_category("Shartnomalar")))
    check("Batafsil kartada kategoriya", "Shartnomalar" in handlers._format_task_detail_card(_ct))
    _imp_cat = handlers._structured_tasks_from_table([("Vazifa", "Kategoriya"), ("Ish", "SMM")])
    check("Import: Kategoriya ustuni o'qiladi", bool(_imp_cat) and _imp_cat[0]["data"].get("category") == "SMM")

    print("\n── Kategoriya boshqaruvi (qo'shish/o'chirish) ──")
    _m1 = await database.create_task({"title": "M1", "priority": "P2", "status": "todo", "category": "TestKat"})
    _m2 = await database.create_task({"title": "M2", "priority": "P2", "status": "todo", "category": "TestKat"})
    check("count_tasks_in_category", await database.count_tasks_in_category("TestKat") == 2)
    await handlers._execute_actions([{"type": "assign_category", "data": {"category": "YangiKat", "from_category": "TestKat"}}])
    check("assign_category (ko'chirish)", await database.count_tasks_in_category("YangiKat") == 2 and await database.count_tasks_in_category("TestKat") == 0)
    await handlers._execute_actions([{"type": "delete_category", "data": {"category": "YangiKat"}}])
    _mt = await database.get_task(_m1)
    check("delete_category: yorliq olinadi, vazifa qoladi", (not _mt.get("category")) and _mt.get("status") == "todo")
    await database.update_task(_m1, {"category": "OchKat"})
    await database.update_task(_m2, {"category": "OchKat"})
    await handlers._execute_actions([{"type": "delete_tasks_by_category", "data": {"category": "OchKat"}}])
    check("delete_tasks_by_category: o'chiriladi", await database.count_tasks_in_category("OchKat") == 0 and (await database.get_task(_m1)) is None)
    check("Kategoriya o'chirish tasdiq to'plamida", {"delete_category", "delete_tasks_by_category"} <= handlers._CATEGORY_DELETE_ACTION_TYPES)
    _pv = await handlers._format_create_preview([{"type": "delete_tasks_by_category", "data": {"category": "Shartnomalar"}}])
    check("Preview: kategoriya o'chirish (son bilan)", "Shartnomalar" in _pv and "o'chiriladi" in _pv)

    print("\n── Kategoriyalar jadvali (B variant: ikonka/arxiv/tartib) ──")
    await database.create_category("Reklama", "🔴")
    await database.create_task({"title": "RT1", "priority": "P2", "status": "todo", "category": "SMM2"})
    _lc = await database.list_categories()
    check("list_categories: bo'sh kategoriya + ikonka", any(c["name"] == "Reklama" and c["count"] == 0 and c["icon"] == "🔴" for c in _lc))
    check("list_categories: orphan ko'rinadi", any(c["name"] == "SMM2" for c in _lc))
    await database.update_category("Reklama", new_name="Reklama2", icon="🟣")
    check("update_category: rename + icon", (lambda g: bool(g) and g["icon"] == "🟣")(await database.get_category("Reklama2")))
    await database.archive_category("Reklama2", True)
    check("archive: faolda yo'q", not any(c["name"] == "Reklama2" for c in await database.list_categories()))
    check("archive: arxivda bor", any(c["name"] == "Reklama2" for c in await database.list_categories(include_archived=True)))
    await database.archive_category("Reklama2", False)
    check("unarchive: faolga qaytadi", any(c["name"] == "Reklama2" for c in await database.list_categories()))
    await database.create_category("MZ1"); await database.create_category("MZ2")
    _b4 = [c["name"] for c in await database.list_categories() if c["name"] in ("MZ1", "MZ2")]
    await database.move_category("MZ2", "up")
    _af = [c["name"] for c in await database.list_categories() if c["name"] in ("MZ1", "MZ2")]
    check("move_category: tartib o'zgaradi", _b4 != _af)
    await handlers._execute_actions([{"type": "create_category", "data": {"category": "ActKat", "icon": "🟢"}}])
    check("action create_category", (await database.get_category("ActKat")) is not None)
    await handlers._execute_actions([{"type": "archive_category", "data": {"category": "ActKat"}}])
    check("action archive_category", (lambda g: bool(g) and g["archived"] == 1)(await database.get_category("ActKat")))
    await database.create_category("DelKat")
    await handlers._execute_actions([{"type": "delete_category", "data": {"category": "DelKat"}}])
    check("action delete_category: metadata o'chadi", (await database.get_category("DelKat")) is None)

    print("\n── Import round-trip dedup (yangilash, dublikat emas) ──")
    _t = await database.create_task({"title": "Dedup sinov", "priority": "P2", "status": "todo"})
    _a2 = handlers._structured_tasks_from_table(
        [("№", "Vazifa", "Ijrochi", "Muddat", "Ustuvorlik", "Holat", "Izoh", "ID"),
         (1, "Dedup YANGI", "Dilshod", "", "Muhim", "Aktiv", "", _t)])
    check("Import: ID ustuni o'qiladi", bool(_a2) and _a2[0].get("_id") == _t)
    for _a in _a2:
        _rid = _a.pop("_id", "")
        if _rid and await database.get_task(_rid):
            _a["type"] = "update_task"; _a["id"] = _rid
    check("Import dedup: mavjud ID → update_task", _a2[0]["type"] == "update_task")
    await handlers._execute_actions(_a2)
    _cnt = len([t for t in await database.list_tasks(limit=999) if t["title"] in ("Dedup sinov", "Dedup YANGI")])
    _upd = await database.get_task(_t)
    check("Import dedup: yangilandi, dublikat YO'Q", _cnt == 1 and _upd["title"] == "Dedup YANGI" and _upd["assignee"] == "Dilshod", f"cnt={_cnt}")

    # Title-based dedup (covers files with NO hidden ID + smart-extracted PDFs)
    _bt = {(t.get("title") or "").strip().lower(): t["id"]
           for t in await database.list_tasks(status_in=["todo", "in_progress", "blocked"], limit=999)
           if t.get("title")}
    _new = [{"type": "create_task", "data": {"title": "Dedup YANGI", "priority": "P0"}},
            {"type": "create_task", "data": {"title": "Mutlaqo yangi XYZ", "priority": "P2"}}]
    _conv = handlers._apply_title_dedup(_new, _bt)
    check("Title dedup: mavjud sarlavha → update", _new[0]["type"] == "update_task" and _conv == 1)
    check("Title dedup: yangi sarlavha → create", _new[1]["type"] == "create_task")

    print("\n── _humanize_error (aniq sabab) ──")
    check("humanize: tarmoq", "ulanish" in handlers._humanize_error(Exception("Cannot connect [Network is unreachable]")).lower())
    check("humanize: bo'sh xabar", "bo'sh xabar" in handlers._humanize_error(Exception("messages.8: user messages must have non-empty content")).lower())
    check("humanize: noma'lum → tur ko'rsatiladi", "ValueError" in handlers._humanize_error(ValueError("nimadir")))

    print("\n── Ikonka birligi (list ↔ batafsil) ──")
    _p2 = {"title": "P2 sinov", "priority": "P2", "status": "todo", "deadline": None}
    check("Badge: P2 list=batafsil bir xil", handlers._format_task_detail_card(_p2).startswith(handlers._task_badge(_p2)))

    print("\n── Delegatsiya → Ijrochilar paneliga birlashtirildi ──")
    try:
        import aiosqlite as _sq
        async with _sq.connect(config.DATABASE_PATH) as _db:
            await _db.execute("SELECT * FROM tasks WHERE status IN ('todo','in_progress') "
                              "AND assignee IS NOT NULL AND LOWER(TRIM(assignee)) NOT IN "
                              "('','men','siz','belgilanmagan','—','oʻzim','o''zim','o''z','ozim') LIMIT 5")
        check("Qotgan-topshiriqlar SQL ishlaydi", True)
    except Exception as e:
        check("Qotgan-topshiriqlar SQL ishlaydi", False, f"{type(e).__name__}: {e}")
    # Birlashtirish: alohida /delegations olib tashlandi, Ijrochilarga ko'chdi
    check("Birlashma: _render_stale_delegations mavjud", hasattr(handlers, "_render_stale_delegations"))
    check("Birlashma: eski cmd_delegations olib tashlandi", not hasattr(handlers, "cmd_delegations"))
    check("Birlashma: '⏳ Kutilayotganlar' tugmasi Ijrochilar klaviaturasida",
          any(b.text == handlers.YBTN_TEAM_STALE
              for row in handlers.team_section_reply_keyboard().keyboard for b in row))
    check("Birlashma: YBTN_TEAM_STALE section-label ro'yxatida",
          handlers.YBTN_TEAM_STALE in handlers._SECTION_LABELS)

    print("\n── Bo'sh slot (free-time) hisoblash ──")
    # SOF funksiya: 11:00-12:30 + 15:00-16:00 band → 3 ta bo'shliq
    d = datetime(2026, 6, 9).date()
    ws = TZ.localize(datetime.combine(d, dtime(9, 0)))
    we = TZ.localize(datetime.combine(d, dtime(18, 0)))
    busy = [(TZ.localize(datetime.combine(d, dtime(11, 0))), TZ.localize(datetime.combine(d, dtime(12, 30)))),
            (TZ.localize(datetime.combine(d, dtime(15, 0))), TZ.localize(datetime.combine(d, dtime(16, 0))))]
    fr = handlers._compute_free_slots(busy, ws, we)
    check("free-slot: 3 ta bo'shliq", len(fr) == 3, f"len={len(fr)}")
    check("free-slot: birinchi 09:00–11:00", fr[0][0].strftime("%H:%M") == "09:00" and fr[0][1].strftime("%H:%M") == "11:00")
    # LMT-bug regressiya: 09:00–11:00 ANIQ 120 daqiqa bo'lishi shart (pytz +04:37 emas)
    check("free-slot: davomiylik aniq (LMT bug yo'q)", int((fr[0][1] - fr[0][0]).total_seconds() // 60) == 120,
          f"{int((fr[0][1]-fr[0][0]).total_seconds()//60)} daq")
    # overlap birlashishi
    ov = [(TZ.localize(datetime.combine(d, dtime(10, 0))), TZ.localize(datetime.combine(d, dtime(12, 0)))),
          (TZ.localize(datetime.combine(d, dtime(11, 0))), TZ.localize(datetime.combine(d, dtime(13, 0))))]
    fov = handlers._compute_free_slots(ov, ws, we)
    check("free-slot: overlap birlashdi", len(fov) == 2 and fov[1][0].strftime("%H:%M") == "13:00")
    # 30 daqiqadan kichik bo'shliq tashlanadi
    tiny = [(ws, TZ.localize(datetime.combine(d, dtime(11, 0)))),
            (TZ.localize(datetime.combine(d, dtime(11, 20))), we)]
    check("free-slot: <30daq bo'shliq tashlandi", handlers._compute_free_slots(tiny, ws, we) == [])
    # bo'sh kun → bitta to'liq slot
    check("free-slot: bo'sh kun = 1 slot", len(handlers._compute_free_slots([], ws, we)) == 1)
    # _fmt_dur
    check("free-slot: _fmt_dur", handlers._fmt_dur(45) == "45 daq" and handlers._fmt_dur(120) == "2 soat" and handlers._fmt_dur(90) == "1s 30daq")
    # _resolve_target_date
    _now = datetime(2026, 6, 6, 10, 0, tzinfo=TZ)  # Shanba
    check("free-slot: sana resolver (ISO + hafta-kun)",
          handlers._resolve_target_date("2026-06-09", _now) == datetime(2026, 6, 9).date()
          and handlers._resolve_target_date("seshanba", _now) == datetime(2026, 6, 9).date()
          and handlers._resolve_target_date("ertaga", _now) == datetime(2026, 6, 7).date())
    check("free-slot: show_free_slots dispatch ro'yxatda", "show_free_slots" in handlers._SHOW_ACTION_TYPES)

    print("\n── Bayonnoma Word dizayni (variant-2: rang-kodli, jadval) ──")
    import io as _io
    from docx import Document as _Docx
    _proto = (
        "📝 **UCHRASHUV BAYONNOMASI**\n━━━━━\n"
        "📅 **Sana va vaqt** — 2026-yil 6-iyun\n"
        "👥 **Ishtirokchilar** — A.Karimov, B.Soliyev\n\n"
        "📋 **KUN TARTIBI**\nBirinchi mavzu.\nIkkinchi mavzu.\n\n"
        "**1.  Birinchi mavzu**\n"
        "ESHITILDI   Masala muhokama qilindi.\n"
        "QAROR   Tasdiqlansin.\n"
        "TOPSHIRIQ   A.Karimov — hisobot tayyorlansin.\n\n"
        "Izoh: 1 ta topshiriq qo'shildi."
    )
    _tasks = [{"assignee": "A.Karimov", "title": "hisobot tayyorlansin.", "deadline": None}]
    _dx = _Docx(_io.BytesIO(handlers._build_protocol_docx_bytes("T", _proto, tasks=_tasks)))

    def _col(r):
        try:
            return str(r.font.color.rgb) if (r.font.color and r.font.color.rgb is not None) else None
        except Exception:
            return None
    _by_kw = {}
    for p in _dx.paragraphs:
        if p.runs:
            _by_kw[p.text.split()[0] if p.text.strip() else ""] = (p.runs[0], p.text)
    check("proto: Arial 13pt Normal", _dx.styles["Normal"].font.name == "Arial" and _dx.styles["Normal"].font.size.pt == 13)
    check("proto: dekorativ sarlavha tashlandi (UCHRASHUV BAYONNOMASI yo'q)",
          not any("UCHRASHUV BAYONNOMASI" in p.text.upper() for p in _dx.paragraphs))
    check("proto: ESHITILDI o'rta yashil (1A8B4D)", "ESHITILDI" in _by_kw and _col(_by_kw["ESHITILDI"][0]) == "1A8B4D")
    check("proto: QAROR to'q yashil (0B6B36)", "QAROR" in _by_kw and _col(_by_kw["QAROR"][0]) == "0B6B36")
    check("proto: TOPSHIRIQ oltin (B8860B)", "TOPSHIRIQ" in _by_kw and _col(_by_kw["TOPSHIRIQ"][0]) == "B8860B")
    _izoh = next((p for p in _dx.paragraphs if p.text.lower().startswith("izoh")), None)
    check("proto: Izoh qizil bold-italic (EE0000)",
          _izoh is not None and _izoh.runs and _col(_izoh.runs[0]) == "EE0000"
          and _izoh.runs[0].bold and _izoh.runs[0].italic)
    _meta = next((p for p in _dx.paragraphs if p.text.startswith("Sana va vaqt")), None)
    check("proto: metadata yorlig'i yashil + tab", _meta is not None and "\t" in _meta.text and _col(_meta.runs[0]) == "0B6B36")
    check("proto: keyword content qora (rangsiz)", "QAROR" in _by_kw and len(_by_kw["QAROR"][0].text) and
          any(_col(r) is None for r in next(p for p in _dx.paragraphs if p.text.startswith("QAROR")).runs))
    check("proto: topshiriqlar jadvali (4 ustun, 1 satr)",
          len(_dx.tables) == 1 and len(_dx.tables[0].columns) == 4 and len(_dx.tables[0].rows) == 2)
    check("proto: jadval sarlavhasi № Mas'ul Topshiriq Muddat",
          [c.text for c in _dx.tables[0].rows[0].cells] == ["№", "Mas'ul", "Topshiriq", "Muddat"])
    check("proto: Muddat null → 'Aniqlashtirilsin'", _dx.tables[0].rows[1].cells[3].text == "Aniqlashtirilsin")

    print("\n── Ikonka nomuvofiqligi tuzatildi (bo'lim ikonkalari noyob) ──")
    import pathlib as _pl
    _src = _pl.Path(handlers.__file__).read_text(encoding="utf-8")
    check("Kategoriyalar reply tugma = 🗄 (🏷 emas)", handlers.TBTN_TASKS_CATEGORIES.startswith("🗄"))
    check("Ijrochilar reply tugma = 👥", handlers.BTN_TEAM.startswith("👥"))
    check("Kategoriya ikonkasi ≠ Ijrochilar ikonkasi",
          handlers.TBTN_TASKS_CATEGORIES[0] != handlers.BTN_TEAM[0])
    check("Kategoriya header = 🗄 (🏷 KATEGORIYALAR yo'q)",
          "🗄 **KATEGORIYALAR**" in _src and "🏷 **KATEGORIYALAR**" not in _src)
    check("Teglar hali 🏷 (kategoriya bilan to'qnashmaydi)", "🏷 Teglar" in _src)
    check("Qotgan-topshiriqlar header = ⏳ (eski DELEGATSIYALAR header yo'q)",
          "⏳  **KUTILAYOTGAN TOPSHIRIQLAR**" in _src and "**DELEGATSIYALAR**" not in _src)
    check("Delegatsiya stats kichik-bo'limi = 📋", "📋 **Delegatsiya**" in _src and "👥 **Delegatsiya**" not in _src)

    print("\n── Qaydlar: sana bo'yicha guruhlash (Variant B) ──")
    _now = datetime.now(handlers.database.TZ)
    _nt = {"id": "a", "title": "Bugungi qayd", "content": "matn bugun",
           "source": "voice", "created_at": _now.isoformat()}
    _ny = {"id": "b", "title": "Kechagi qayd", "content": "matn kecha",
           "source": "forward", "created_at": (_now - timedelta(days=1)).isoformat()}
    _nout = handlers._format_notes_compact([_nt, _ny], "inbox", inbox_count=2)
    check("Qaydlar header = 'QAYDLAR · INBOX · 2 ta'", "QAYDLAR · INBOX · 2 ta" in _nout)
    check("Qaydlar: 📅 BUGUN guruhi", "📅 **BUGUN**" in _nout)
    check("Qaydlar: 📅 KECHA guruhi", "📅 **KECHA**" in _nout)
    check("Qaydlar: manba ikonkasi sarlavha oldida (🎙 1.)", "🎙 1." in _nout)
    check("Qaydlar: forward ikonkasi (🔁 2.)", "🔁 2." in _nout)
    check("Qaydlar: eski 'NOTES ·' sarlavhasi yo'q", "NOTES ·" not in _nout)

    print("\n── Qaydlarga junk tushmasligi (tugma/buyruq/bot-chiqishi filtri) ──")
    check("noise: '/team' (buyruq) bloklanadi", handlers._is_note_noise("/team"))
    check("noise: '⬅️ Asosiy menyu' (tugma) bloklanadi", handlers._is_note_noise("⬅️ Asosiy menyu"))
    check("noise: real matn bloklanmaydi", not handlers._is_note_noise("Pepsi bilan shartnoma tuzish"))
    _botout = "📌  VAZIFALAR\n\nKo'rinish · Aktiv\nNatija · 9 ta\n━━━━━━━━━━━━━━━"
    check("bot-echo: botning '📌 VAZIFALAR' paneli aniqlanadi", handlers._looks_like_bot_output(_botout))
    check("bot-echo: real matn (divider yo'q) aniqlanmaydi",
          not handlers._looks_like_bot_output("Ertaga soat 10da uchrashuv bor"))
    check("forward guard: _forward_is_bot_echo mavjud", hasattr(handlers, "_forward_is_bot_echo"))

    print("\n── Qayd detali: tugma to'qnashuvi + ixcham karta ──")
    _dnote = {"id": "n-x", "title": "Test qayd", "content": "matn",
              "source": "llm", "created_at": "2026-06-06T02:41:00+05:00",
              "status": "inbox", "tags": ["bot", "UI"]}
    _dkb = [b.text for row in handlers.note_detail_menu(_dnote).inline_keyboard for b in row]
    check("dedup: inline action = '📦 Arxivga' (yo'nalishli)", "📦 Arxivga" in _dkb)
    check("dedup: inline'da '📦 Arxiv' (filtr nomi) YO'Q", "📦 Arxiv" not in _dkb)
    check("dedup: reply filtr hali '📦 Arxiv'", handlers.NBTN_NOTES_ARCHIVED == "📦 Arxiv")
    check("dedup: action ≠ filtr (to'qnashuv yo'q)",
          "📦 Arxivga" != handlers.NBTN_NOTES_ARCHIVED)
    _dtext, _ = handlers._format_note_detail(_dnote)
    check("karta: meta bitta qatorda (manba · sana · holat)", "🤖 LLM · 📅 06-06 02:41 · 📥 Inbox" in _dtext)
    check("karta: ortiqcha '🔖 Holat:' qatori olib tashlandi", "🔖 Holat:" not in _dtext)
    check("karta: teglar bo'sh joy bilan (vergulsiz)", "#bot #UI" in _dtext)

    print("\n── Inline ↔ Reply takror yo'q (konsepsiya: filtr=reply, inline=kontent) ──")
    _rnotes = [{"id": f"n-{i}", "title": f"Q{i}", "content": "x", "source": "manual",
                "created_at": "2026-06-06T10:00:00+05:00"} for i in range(12)]
    _list_inline = {b.text for row in handlers.notes_compact_keyboard(_rnotes, "inbox", 1).inline_keyboard for b in row}
    _reply = {b.text for row in handlers.notes_section_reply_keyboard().keyboard for b in row}
    _det_inline = {b.text for row in handlers.note_detail_menu(_dnote).inline_keyboard for b in row}
    check("list inline'da filtr-pills YO'Q", not ({"📥 Inbox", "⚙️ Ishlangan", "📦 Arxiv"} & _list_inline))
    check("list inline ∩ reply = bo'sh", not (_list_inline & _reply), str(_list_inline & _reply))
    check("detail inline ∩ reply = bo'sh", not (_det_inline & _reply), str(_det_inline & _reply))
    check("list inline'da pagination saqlangan", "➡️" in _list_inline)
    check("list inline'da drill-down raqamlar bor", "1" in _list_inline)

    print("\n── Qaydlar takomillashtirish (A · B3 Ishlandi · B4 yosh · C6 split) ──")
    check("A2: 'Yangi qayd' tugmasi (note emas)", handlers.NBTN_NOTES_NEW == "➕ Yangi qayd")
    check("A2: 'Qayd qidirish' tugmasi", handlers.NBTN_NOTES_SEARCH == "🔍 Qayd qidirish")
    _nb = {b.text for row in handlers.notes_section_reply_keyboard().keyboard for b in row}
    check("A2: reply tugmalarда inglizcha 'note' yo'q", not any("note" in b.lower() for b in _nb))
    check("A1: diagnostika header 🩺 (🔍 emas)", "🩺  **DIAGNOSTIKA**" in _src)
    check("B3: database.mark_note_done mavjud", hasattr(database, "mark_note_done"))
    _now2 = datetime.now(handlers.database.TZ)
    _inb = {"id": "x", "status": "inbox", "title": "t", "content": "x",
            "source": "manual", "created_at": _now2.isoformat(), "tags": []}
    _im = [b.text for row in handlers.note_detail_menu(_inb).inline_keyboard for b in row]
    check("B3: inbox menyuда '✅ Ishlandi'", "✅ Ishlandi" in _im)
    _proc = dict(_inb); _proc["status"] = "processed"
    _pm = [b.text for row in handlers.note_detail_menu(_proc).inline_keyboard for b in row]
    check("B3: processed menyuда Ishlandi/ajrat YO'Q",
          "✅ Ishlandi" not in _pm and not any("ajrat" in x for x in _pm))
    _old = dict(_inb); _old["created_at"] = (_now2 - timedelta(days=4)).isoformat()
    check("B4: eski inbox qayd yoshini ko'rsatadi (4 kun)", "4 kun" in handlers._format_note_detail(_old)[0])
    check("B4: bugungi qaydда yosh yo'q",
          "kun" not in handlers._format_note_detail(_inb)[0].split("blockquote")[0])
    check("C6: inbox menyuда '✂️ ...ajrat'", any("ajrat" in x for x in _im))
    check("C6: split directive create_task so'raydi",
          "create_task" in handlers._build_note_split_directive("a, b, c"))

    print("\n── Eslatmalar bo'limi (Bugun→eslatma · Barchasi📋 · Keyingi olib · edit) ──")
    _rk = {b.text for row in handlers.reminders_section_reply_keyboard().keyboard for b in row}
    check("Eslatma: 'Keyingi' tugmasi olib tashlandi", not any("Keyingi" in b for b in _rk))
    check("Eslatma: 'Barchasi' = 📋 (tasks bilan moslik)", "📋 Barchasi" in _rk)
    check("Eslatma: filtrlar today/sent/all (upcoming yo'q)",
          set(handlers._REMINDERS_SECTION_FILTERS.values()) == {"today", "sent", "all"})
    check("Eslatma: ⏰ Bugun → 'today'", handlers._REMINDERS_SECTION_FILTERS.get(handlers.RBTN_REMINDERS_TODAY) == "today")
    check("Eslatma: Bugun state-mustaqil handleri bor (tasks leak yo'q)",
          hasattr(handlers, "handle_reminder_filter_anystate"))
    _rdm = {b.text for row in handlers.reminder_detail_menu({"id": "r", "status": "scheduled"}).inline_keyboard for b in row}
    check("Eslatma: '✏️ Tahrirlash' tugmasi mavjud (konsolidatsiya)", "✏️ Tahrirlash" in _rdm)
    check("Eslatma: Bajarildi/o'chir tugmalari", "✅ Bajarildi" in _rdm and "🗑 O'chirish" in _rdm)

    print("\n── Professional: lock · backup/undo · onboarding · audit · auto-chase ──")
    import pathlib as _pl2
    import bot as _bot
    import scheduler as _sch
    check("A1: single-instance lock funksiyasi", hasattr(_bot, "_acquire_single_instance_lock"))
    _svc = _pl2.Path(handlers.__file__).parent / "deploy" / "yordamchi.service"
    check("A2: systemd xizmat fayli mavjud", _svc.exists())
    check("A2: xizmat Restart=always", _svc.exists() and "Restart=always" in _svc.read_text(encoding="utf-8"))
    check("B: _create_db_backup mavjud", hasattr(handlers, "_create_db_backup"))
    check("B: cb_undo_delete + _UNDO_BACKUPS",
          hasattr(handlers, "cb_undo_delete") and hasattr(handlers, "_UNDO_BACKUPS"))
    check("C1: /start onboarding (imkoniyatlar tavsifi)", "Imkoniyatlar" in _src)
    check("C2: database.list_recent_actions", hasattr(database, "list_recent_actions"))
    check("C2: diagnostikada 'So'nggi amallar' bo'limi", "So'nggi amallar" in _src)
    check("D: database.list_stale_delegations", hasattr(database, "list_stale_delegations"))
    check("D: scheduler._stale_delegation_digest", hasattr(_sch.YordamchiScheduler, "_stale_delegation_digest"))

    print("\n── Bayonnomalar markaziy ro'yxati (oy bo'yicha) ──")
    check("Bayonnoma: database.list_meetings_with_protocol", hasattr(database, "list_meetings_with_protocol"))
    check("Bayonnoma: heuristika — uzun matn = protokol",
          handlers._looks_like_protocol(["📝 UCHRASHUV BAYONNOMASI uzun matn shu yerda joylashgan"]))
    check("Bayonnoma: heuristika — task-id ro'yxat ≠ protokol",
          not handlers._looks_like_protocol(["t-123", "t-456"]))
    check("Bayonnoma: heuristika — bo'sh ≠ protokol", not handlers._looks_like_protocol([]))
    check("Bayonnoma: _render_protocols + cmd_protocols",
          hasattr(handlers, "_render_protocols") and hasattr(handlers, "cmd_protocols"))
    _mk = {b.text for row in handlers.meetings_section_reply_keyboard().keyboard for b in row}
    check("Bayonnoma: '📄 Bayonnomalar' tugmasi meetings kbd'da", "📄 Bayonnomalar" in _mk)
    # 'ulashish ishlamadi' bug: share handler state'ga emas, uchrashuvga ham tayanishi shart
    import re as _re3
    _share_fn = _re3.search(
        r"async def cb_protocol_share\(.*?(?=\n@router|\nasync def |\ndef )", _src, _re3.DOTALL)
    check("Bayonnoma: Ulashish uchrashuvdan fallback o'qiydi (bug tuzatildi)",
          bool(_share_fn) and "get_meeting" in _share_fn.group(0)
          and "follow_up_actions" in _share_fn.group(0))
    # Inline ulashish (bayonnoma → istalgan chatga)
    check("Inline: yagona handle_inline_query mavjud", hasattr(handlers, "handle_inline_query"))
    check("Inline: aynan BITTA @router.inline_query() handler (dublikat yo'q — leak bug)",
          len(handlers.router.inline_query.handlers) == 1)
    check("Inline: 'Ulashish' switch_inline_query mavjud (📋 Nusxa = proto_share callback)",
          'switch_inline_query=f"proto:{mid}"' in _src and 'callback_data=f"proto_share:{mid}"' in _src)

    class _IqU:
        def __init__(s, i): s.id = i
    class _IqQ:
        def __init__(s, uid, q): s.from_user = _IqU(uid); s.query = q; s.res = None
        async def answer(s, results, **k): s.res = results
    _imid = await database.create_meeting({
        "title": "InlineTest", "datetime_start": "2026-06-06T10:00:00+05:00",
        "participants": [], "follow_up_actions": ["Bu bayonnoma matni ulashish uchun."]})
    _iq = _IqQ(config.PRINCIPAL_USER_ID, f"proto:{_imid}")
    await handlers.handle_inline_query(_iq)
    check("Inline: principal bayonnomani oladi", bool(_iq.res) and len(_iq.res) == 1)
    _iq2 = _IqQ(999999999, f"proto:{_imid}")
    await handlers.handle_inline_query(_iq2)
    check("Inline: begona foydalanuvchi natija olmaydi (gate)", _iq2.res == [])
    _iq3 = _IqQ(config.PRINCIPAL_USER_ID, "proto:m-DOESNOTEXIST")
    await handlers.handle_inline_query(_iq3)
    check("Inline: yo'q bayonnoma → xom 'proto:' matni QAYTMAYDI (leak bug tuzatildi)",
          bool(_iq3.res) and all(
              "proto:" not in (r.input_message_content.message_text or "") for r in _iq3.res))
    import aiosqlite as _sq5
    async with _sq5.connect(config.DATABASE_PATH) as _db5:
        await _db5.execute("DELETE FROM meetings WHERE id=?", (_imid,)); await _db5.commit()

    print("\n── 2 xil ulashish: 📋 Nusxa (qo'lda) + 📤 Ulashish (inline) ──")
    _prk = handlers._protocol_result_kb("m1", 1, saved=True, tasks_done=False)
    _prb = [(b.text, bool(b.switch_inline_query)) for row in _prk.inline_keyboard for b in row]
    check("Bayonnoma: 📋 Nusxa (qo'lda) + 📤 Ulashish (inline)",
          ("📋 Nusxa", False) in _prb and ("📤 Ulashish", True) in _prb)
    _pol = handlers._build_keyboard(
        [[{"label": "📋 Nusxa olish", "callback": "copy"},
          {"label": "x", "callback": "share"},
          {"label": "✎", "callback": "edit:polish"}]],
        {}, share_text="**Tahrirlangan matn:**\n───\nSalom hamkor.\n───")
    _pb = [(b.text, bool(b.switch_inline_query)) for row in _pol.inline_keyboard for b in row]
    check("Sayqal matn: 📋 Nusxa + 📤 Ulashish (inline), no-op copy tashlandi",
          ("📋 Nusxa", False) in _pb and ("📤 Ulashish", True) in _pb
          and not any(t == "📋 Nusxa olish" for t, _ in _pb))
    _itok = next((b.switch_inline_query.split("txt:")[1]
                  for row in _pol.inline_keyboard for b in row
                  if b.switch_inline_query and "txt:" in b.switch_inline_query), None)
    check("Inline kesh DB-backed (restart'ga bardoshli, toza matn)",
          _itok is not None and (await database.get_share_text(_itok)) == "Salom hamkor.")

    print("\n── Eslatma detal tugmalari (kontekstga mos) ──")
    def _rmenu(r):
        return [b.text for row in handlers.reminder_detail_menu(r).inline_keyboard for b in row]
    _rec = _rmenu({"id": "r", "status": "scheduled", "recurrence_rule": "weekly",
                   "remind_at": "2026-06-08T12:00:00+05:00"})
    _one = _rmenu({"id": "r", "status": "scheduled", "recurrence_rule": None,
                   "remind_at": "2026-06-08T12:00:00+05:00"})
    _dn = _rmenu({"id": "r", "status": "done", "recurrence_rule": None})
    check("Eslatma: takroriyda ⏭ skip + 🛑 stop",
          any("o'tkaz" in x for x in _rec) and any("to'xtat" in x.lower() for x in _rec))
    check("Eslatma: takroriyda 'Bajarildi' YO'Q (seriya o'chmaydi)", "✅ Bajarildi" not in _rec)
    check("Eslatma: bir martalikda ✅ Bajarildi + snooze",
          "✅ Bajarildi" in _one and "⏰ 15 daq" in _one)
    check("Eslatma: snooze hammasi ⏰ (izchil)",
          all(x.startswith("⏰") for x in _one if ("daq" in x or "soat" in x or "Ertaga" in x)))
    check("Eslatma: bir martalik ham '✏️ Tahrirlash' (eski 📆/📅 inline yo'q)",
          "✏️ Tahrirlash" in _one and "📆 Vaqt" not in str(_one + _rec))
    check("Eslatma: done → ↺ Qayta eslat", any("Qayta eslat" in x for x in _dn))
    check("Eslatma: cb_reminder_skip + cb_reminder_stop",
          hasattr(handlers, "cb_reminder_skip") and hasattr(handlers, "cb_reminder_stop"))
    check("Eslatma: done'da ham '✏️ Tahrirlash'", "✏️ Tahrirlash" in _dn)
    check("Eslatma: default ko'rinish 'active' (done emas, aktiv chiqadi)",
          '_render_reminders_for_filter(message, "active")' in _src)
    _act, _albl = await handlers._load_reminders_for_filter("active")
    check("Eslatma: 'active' filtr done'ni ko'rsatmaydi",
          all(r.get("status") != "done" for r in _act))
    check("Eslatma: Barchasi tugmasi hali 'all' (done ko'rinadi)",
          handlers._REMINDERS_SECTION_FILTERS.get(handlers.RBTN_REMINDERS_ALL) == "all")

    print("\n── Eslatma tahrir menyusi (konsolidatsiya + Takror + Izoh + weekdays) ──")
    _em = [b.text for row in handlers.reminder_detail_menu(
        {"id": "r", "status": "scheduled", "recurrence_rule": "weekly"}).inline_keyboard for b in row]
    check("Tahrir: bitta '✏️ Tahrirlash' (Matn/Vaqt alohida emas)",
          "✏️ Tahrirlash" in _em and "✏️ Matn" not in _em)
    check("Tahrir: edit-menu/recurrence callbacklari mavjud",
          hasattr(handlers, "cb_reminder_edit_menu") and hasattr(handlers, "cb_reminder_recurrence_menu")
          and hasattr(handlers, "cb_reminder_set_recurrence"))
    check("Takror: variantlarda 'weekdays' (ish kunlari) bor",
          "weekdays" in [v for v, _ in handlers._RECUR_OPTIONS])
    check("weekdays: normalize ('ish kunlari'/'Dushanba-juma')",
          database.normalize_recurrence_rule("ish kunlari") == "weekdays"
          and database.normalize_recurrence_rule("Dushanba-juma") == "weekdays")
    check("weekdays: label 'ish kunlari'", "ish kunlari" in handlers._format_recurrence_label("weekdays"))
    check("weekdays: dam olishni o'tkazadi (Juma 12-iyun → Dushanba 15-iyun)",
          database.compute_next_recurrence("2026-06-12T12:00:00+05:00", "weekdays")[:10] == "2026-06-15")

    print("\n── Eslatma tahrir: matn + OVOZ ishlaydi (bug tuzatildi) ──")
    class _ReFakeState:
        def __init__(s, d): s._d = d
        async def get_data(s): return s._d
        async def clear(s): pass
    class _ReFakeMsg:
        def __init__(s, text=None, voice=False):
            s.text = text; s.voice = (object() if voice else None); s.bot = None
            s.chat = type("C", (), {"id": 1})()
        async def answer(s, *a, **k): pass
    _orig_tt = handlers._get_text_or_transcribe
    try:
        _rid = await database.create_reminder(
            {"title": "Eski", "remind_at": "2026-06-10T10:00:00+05:00", "status": "scheduled"})
        async def _vt(m, bot=None): return "Ovozli yangi sarlavha"
        handlers._get_text_or_transcribe = _vt
        await handlers.handle_reminder_edit_value(
            _ReFakeMsg(text=None, voice=True), _ReFakeState({"reminder_id": _rid, "field": "title"}))
        _r = await database.get_reminder(_rid)
        check("Tahrir: OVOZ orqali sarlavha saqlanadi", bool(_r) and _r["title"] == "Ovozli yangi sarlavha")
        async def _vt2(m, bot=None): return "ertaga 09:00"
        handlers._get_text_or_transcribe = _vt2
        await handlers.handle_reminder_edit_value(
            _ReFakeMsg(text=None, voice=True), _ReFakeState({"reminder_id": _rid, "field": "time"}))
        _r = await database.get_reminder(_rid)
        check("Tahrir: OVOZ orqali vaqt saqlanadi", bool(_r) and _r["remind_at"][:10] != "2026-06-10")
        import aiosqlite as _sq3
        async with _sq3.connect(config.DATABASE_PATH) as _db3:
            await _db3.execute("DELETE FROM reminders WHERE id=?", (_rid,)); await _db3.commit()
    finally:
        handlers._get_text_or_transcribe = _orig_tt
    # Kengroq ovoz bug: section handlerlari ovozni transkripsiya qilib, keyin
    # message.text (ovozda None) ishlatardi → crash / yo'qolish. Endi _msg_text.
    check("Ovoz: 'label = message.text' qolmadi (section crash yo'q)",
          "label = message.text.strip()" not in _src)
    check("Ovoz: '_process_and_reply(message, message.text' qolmadi (ovoz Claude'ga)",
          "_process_and_reply(message, message.text" not in _src)

    print("\n" + "=" * 48)
    print(f"NATIJA:  ✅ {PASS} o'tdi   ❌ {FAIL} yiqildi")
    if FAILED:
        print("Yiqilganlar:", ", ".join(FAILED))
    print("=" * 48)
    # cleanup temp db
    try:
        os.remove(_TMP)
    except OSError:
        pass
    return 0 if FAIL == 0 else 1


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
